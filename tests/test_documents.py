import io
import shutil
import subprocess
from pathlib import Path

import pytest
from docx import Document
from reportlab.pdfgen.canvas import Canvas

from agentic_rag.documents import DocumentError, chunk_document, load_document


def test_text_chunks_overlap_and_stable_ids() -> None:
    content = b"abcdefghij"
    chunks = chunk_document("../notes.txt", content, chunk_size=6, overlap=2)
    assert [c.text for c in chunks] == ["abcdef", "efghij"]
    assert all(c.source == "notes.txt" for c in chunks)
    assert chunks == chunk_document("../notes.txt", content, chunk_size=6, overlap=2)
    assert chunks[0].id != chunk_document("notes.txt", b"changed")[0].id


def test_pdf_page_numbers() -> None:
    output = io.BytesIO()
    canvas = Canvas(output)
    canvas.drawString(50, 700, "Refunds take five days.")
    canvas.showPage()
    canvas.drawString(50, 700, "Second page content.")
    canvas.save()
    chunks = chunk_document("policy.pdf", output.getvalue())
    assert [c.page for c in chunks] == [1, 2]
    assert "five days" in chunks[0].text


def test_docx_paragraphs_and_tables() -> None:
    output = io.BytesIO()
    doc = Document()
    doc.add_paragraph("Atlas policy")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Refund"
    table.cell(0, 1).text = "30 days"
    doc.save(output)
    assert "Refund | 30 days" in load_document("policy.docx", output.getvalue())[0].text


def test_legacy_doc_uses_safe_subprocess(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr("agentic_rag.documents.shutil.which", lambda _: "/usr/bin/antiword")

    def fake_run(args: list[str], **kwargs: object) -> subprocess.CompletedProcess[bytes]:
        assert args[:3] == ["/usr/bin/antiword", "-m", "UTF-8.txt"]
        assert kwargs["timeout"] == 20
        assert "shell" not in kwargs
        return subprocess.CompletedProcess(args, 0, b"Legacy document text", b"")

    monkeypatch.setattr("agentic_rag.documents.subprocess.run", fake_run)
    assert load_document("untrusted.doc", b"sample")[0].text == "Legacy document text"


def test_missing_antiword(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr("agentic_rag.documents.shutil.which", lambda _: None)
    with pytest.raises(DocumentError, match="antiword"):
        load_document("policy.doc", b"sample")


@pytest.mark.parametrize(
    ("name", "content"),
    [("a.txt", b""), ("a.txt", b"   "), ("a.exe", b"a"), ("a.pdf", b"broken")],
)
def test_invalid_documents(name: str, content: bytes) -> None:
    with pytest.raises(DocumentError):
        load_document(name, content)


def test_bad_chunk_parameters() -> None:
    with pytest.raises(ValueError):
        chunk_document("a.txt", b"a", chunk_size=10, overlap=10)


@pytest.mark.skipif(
    shutil.which("antiword") is None, reason="antiword system package not installed"
)
def test_real_legacy_doc_fixture() -> None:
    content = (Path(__file__).parent / "fixtures/atlas_legacy.doc").read_bytes()
    assert "30 days" in load_document("atlas_legacy.doc", content)[0].text
