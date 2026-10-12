import hashlib
import io
import shutil
import subprocess
import tempfile
import zipfile
from dataclasses import dataclass
from pathlib import Path

from docx import Document
from pypdf import PdfReader

MAX_FILE_BYTES = 20 * 1024 * 1024
MAX_EXPANDED_BYTES = 50 * 1024 * 1024
SUPPORTED_EXTENSIONS = {".pdf", ".doc", ".docx", ".txt"}


class DocumentError(ValueError):
    pass


@dataclass(frozen=True)
class Page:
    text: str
    number: int


@dataclass(frozen=True)
class Chunk:
    id: str
    document_id: str
    source: str
    page: int
    index: int
    text: str


def load_document(name: str, content: bytes) -> list[Page]:
    suffix = Path(name).suffix.lower()
    if suffix not in SUPPORTED_EXTENSIONS:
        raise DocumentError("Supported formats: PDF, DOC, DOCX and UTF-8 TXT.")
    if not content:
        raise DocumentError("The document is empty.")
    if len(content) > MAX_FILE_BYTES:
        raise DocumentError("Documents must be 20 MB or smaller.")
    try:
        if suffix == ".pdf":
            reader = PdfReader(io.BytesIO(content))
            if reader.is_encrypted:
                raise DocumentError("Password-protected PDFs are not supported.")
            pages = [Page(page.extract_text() or "", i + 1) for i, page in enumerate(reader.pages)]
        elif suffix == ".docx":
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                if sum(item.file_size for item in archive.infolist()) > MAX_EXPANDED_BYTES:
                    raise DocumentError("Expanded DOCX exceeds the 50 MB safety limit.")
            document = Document(io.BytesIO(content))
            paragraphs = [p.text for p in document.paragraphs]
            tables = [
                " | ".join(c.text for c in row.cells) for t in document.tables for row in t.rows
            ]
            pages = [Page("\n".join(paragraphs + tables), 1)]
        elif suffix == ".doc":
            executable = shutil.which("antiword")
            if executable is None:
                raise DocumentError(
                    "Legacy DOC needs antiword. Install it or save the file as DOCX."
                )
            with tempfile.TemporaryDirectory(prefix="rag-doc-") as directory:
                path = Path(directory) / "upload.doc"
                path.write_bytes(content)
                result = subprocess.run(
                    [executable, "-m", "UTF-8.txt", str(path)],
                    capture_output=True,
                    timeout=20,
                    check=False,
                )
            if result.returncode:
                raise DocumentError("Could not read this DOC. Try saving it as DOCX.")
            pages = [Page(result.stdout.decode("utf-8", errors="replace"), 1)]
        else:
            pages = [Page(content.decode("utf-8-sig"), 1)]
    except DocumentError:
        raise
    except Exception as error:
        raise DocumentError(
            f"Could not read {suffix} document. Check its format/encoding."
        ) from error
    pages = [Page(p.text.replace("\x00", "").strip(), p.number) for p in pages if p.text.strip()]
    if not pages:
        raise DocumentError("No extractable text found. Scanned PDFs need OCR before upload.")
    return pages


def chunk_document(
    name: str, content: bytes, *, chunk_size: int = 1800, overlap: int = 200
) -> list[Chunk]:
    if chunk_size <= 0 or overlap < 0 or overlap >= chunk_size:
        raise ValueError("Require chunk_size > overlap >= 0.")
    source = name.replace("\\", "/").rsplit("/", 1)[-1]
    document_id = hashlib.sha256(source.encode() + b"\0" + content).hexdigest()
    chunks: list[Chunk] = []
    for page in load_document(name, content):
        for start in range(0, len(page.text), chunk_size - overlap):
            text = page.text[start : start + chunk_size].strip()
            if text:
                index = len(chunks)
                chunk_id = hashlib.sha256(
                    f"{document_id}:{page.number}:{index}".encode()
                ).hexdigest()
                chunks.append(Chunk(chunk_id, document_id, source, page.number, index, text))
            if start + chunk_size >= len(page.text):
                break
    return chunks
